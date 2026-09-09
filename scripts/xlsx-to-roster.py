#!/usr/bin/env python3
"""권한 체크표(xlsx) → bulk-users.mjs 가 읽는 roster.csv"""
import sys
from openpyxl import load_workbook

TABS = [('sunday','주일예배'), ('youth','청년교구'), ('school','주일학교'),
        ('wd','주중·새벽'), ('district','교구'), ('newfam','새가족'),
        ('special','특별예배'), ('yearcomp','연도비교')]

def main(src, dst):
    ws = load_workbook(src, data_only=True)['계정']
    head = [str(c.value or '').strip() for c in ws[1]]
    idx = {h: i for i, h in enumerate(head)}
    out = ['# 권한 체크표에서 자동 생성. 직접 고치지 말고 xlsx 를 고치세요.',
           '아이디,이름,역할,탭,주일학교부서,청년부,담당팀,교구,초기비밀번호']
    checked = lambda v: str(v or '').strip().upper() in ('O', 'V', 'X', 'Y', '1', 'TRUE', '예')
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        get = lambda h: row[idx[h]] if h in idx and idx[h] < len(row) else None
        uid = str(get('아이디') or '').strip()
        if not uid:
            continue
        admin = checked(get('관리자'))
        tabs = 'home/' + '/'.join(k for k, label in TABS if checked(get(label)))
        cell = lambda h: str(get(h) or '').strip().replace(',', '/')
        pw = str(get('초기비밀번호') or '').strip()
        out.append(','.join([uid, str(get('이름') or '').strip(),
                             'admin' if admin else 'staff', '' if admin else tabs,
                             cell('주일학교 담당부서'), cell('담당 청년부'),
                             cell('담당 팀'), cell('담당 교구'), pw]))
        n += 1
    open(dst, 'w', encoding='utf-8').write('\n'.join(out) + '\n')
    print(f'{dst}  —  {n}명')

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
